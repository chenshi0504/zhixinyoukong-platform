"""
文档加载器模块
支持多种格式的文档解析和内容提取
"""

import os
import re
import json
from pathlib import Path
from typing import List, Dict, Optional
from datetime import datetime
import logging

logger = logging.getLogger(__name__)


class Document:
    """文档对象"""
    def __init__(self, content: str, metadata: Dict):
        self.content = content
        self.metadata = metadata
        self.page_content = content  # LangChain兼容


class MarkdownLoader:
    """Markdown文档加载器"""

    def __init__(self, file_path: str):
        self.file_path = file_path

    def load(self) -> List[Document]:
        """加载Markdown文档"""
        try:
            with open(self.file_path, 'r', encoding='utf-8') as f:
                content = f.read()

            # 提取元数据
            metadata = {
                'source': self.file_path,
                'type': 'markdown',
                'filename': os.path.basename(self.file_path)
            }

            # 提取标题作为额外元数据
            title_match = re.search(r'^#\s+(.+)$', content, re.MULTILINE)
            if title_match:
                metadata['title'] = title_match.group(1)

            return [Document(content=content, metadata=metadata)]

        except Exception as e:
            logger.error(f"加载Markdown文件失败 {self.file_path}: {e}")
            return []


class PythonLoader:
    """Python代码加载器"""

    def __init__(self, file_path: str):
        self.file_path = file_path

    def load(self) -> List[Document]:
        """加载Python文件，提取文档字符串和注释"""
        try:
            with open(self.file_path, 'r', encoding='utf-8') as f:
                content = f.read()

            # 提取模块级文档字符串
            docstring_match = re.search(r'^"""(.*?)"""', content, re.DOTALL | re.MULTILINE)
            module_doc = docstring_match.group(1).strip() if docstring_match else ""

            # 提取函数和类的文档字符串
            functions = re.findall(r'def\s+(\w+)\s*\([^)]*\):\s*"""(.*?)"""', content, re.DOTALL)
            classes = re.findall(r'class\s+(\w+).*?:\s*"""(.*?)"""', content, re.DOTALL)

            # 构建文档内容
            doc_content = f"# {os.path.basename(self.file_path)}\n\n"

            if module_doc:
                doc_content += f"## 模块说明\n{module_doc}\n\n"

            if classes:
                doc_content += "## 类定义\n\n"
                for class_name, class_doc in classes:
                    doc_content += f"### {class_name}\n{class_doc.strip()}\n\n"

            if functions:
                doc_content += "## 函数定义\n\n"
                for func_name, func_doc in functions:
                    doc_content += f"### {func_name}\n{func_doc.strip()}\n\n"

            metadata = {
                'source': self.file_path,
                'type': 'python',
                'filename': os.path.basename(self.file_path),
                'functions': [f[0] for f in functions],
                'classes': [c[0] for c in classes]
            }

            return [Document(content=doc_content, metadata=metadata)]

        except Exception as e:
            logger.error(f"加载Python文件失败 {self.file_path}: {e}")
            return []


class JavaScriptLoader:
    """JavaScript代码加载器"""

    def __init__(self, file_path: str):
        self.file_path = file_path

    def load(self) -> List[Document]:
        """加载JavaScript文件，提取JSDoc注释"""
        try:
            with open(self.file_path, 'r', encoding='utf-8') as f:
                content = f.read()

            # 提取JSDoc注释
            jsdoc_pattern = r'/\*\*(.*?)\*/'
            jsdocs = re.findall(jsdoc_pattern, content, re.DOTALL)

            # 提取函数定义
            func_pattern = r'(?:function|const|let|var)\s+(\w+)\s*[=\(]'
            functions = re.findall(func_pattern, content)

            # 构建文档内容
            doc_content = f"# {os.path.basename(self.file_path)}\n\n"

            if jsdocs:
                doc_content += "## 文档注释\n\n"
                for doc in jsdocs:
                    doc_content += f"{doc.strip()}\n\n"

            metadata = {
                'source': self.file_path,
                'type': 'javascript',
                'filename': os.path.basename(self.file_path),
                'functions': functions
            }

            return [Document(content=doc_content, metadata=metadata)]

        except Exception as e:
            logger.error(f"加载JavaScript文件失败 {self.file_path}: {e}")
            return []


class DocxLoader:
    """Word文档加载器"""

    def __init__(self, file_path: str):
        self.file_path = file_path

    def load(self) -> List[Document]:
        """加载Word文档"""
        try:
            from docx import Document as DocxDocument

            doc = DocxDocument(self.file_path)

            # 提取所有段落
            content = "\n\n".join([para.text for para in doc.paragraphs if para.text.strip()])

            metadata = {
                'source': self.file_path,
                'type': 'docx',
                'filename': os.path.basename(self.file_path)
            }

            return [Document(content=content, metadata=metadata)]

        except ImportError:
            logger.error("需要安装python-docx: pip install python-docx")
            return []
        except Exception as e:
            logger.error(f"加载Word文件失败 {self.file_path}: {e}")
            return []


class PDFLoader:
    """PDF文档加载器"""

    def __init__(self, file_path: str):
        self.file_path = file_path

    def load(self) -> List[Document]:
        """加载PDF文档"""
        try:
            import PyPDF2

            content = ""
            with open(self.file_path, 'rb') as f:
                pdf_reader = PyPDF2.PdfReader(f)
                num_pages = len(pdf_reader.pages)

                for page_num in range(num_pages):
                    page = pdf_reader.pages[page_num]
                    content += page.extract_text() + "\n\n"

            metadata = {
                'source': self.file_path,
                'type': 'pdf',
                'filename': os.path.basename(self.file_path),
                'num_pages': num_pages
            }

            return [Document(content=content.strip(), metadata=metadata)]

        except ImportError:
            logger.error("需要安装PyPDF2: pip install PyPDF2")
            return []
        except Exception as e:
            logger.error(f"加载PDF文件失败 {self.file_path}: {e}")
            return []


class ExperimentResultLoader:
    """实验结果JSON文件加载器"""

    # 模块ID到中文名称的映射
    MODULE_NAMES = {
        'data_collection': '数据采集',
        'data_fusion': '数据融合',
        'traffic_reconstruction': '交通流重构',
        'traffic_prediction': '交通流预测',
        'signal_optimization': '信号优化',
        'path_guidance': '路径诱导',
        'speed_control': '车速控制',
        'vehicle_speed_control': '车速控制',
        'three_level_control': '三层级协同控制',
        'cooperative_control': '协同控制',
        'hierarchical_control': '分层控制',
        'carla_simulation': 'CARLA仿真',
    }

    def __init__(self, file_path: str):
        self.file_path = file_path

    def load(self) -> List[Document]:
        """加载实验结果JSON文件"""
        try:
            with open(self.file_path, 'r', encoding='utf-8') as f:
                data = json.load(f)

            # 提取元数据
            experiment_id = data.get('experiment_id', 'unknown')
            module_id = data.get('module_id', 'unknown')
            module_name = self.MODULE_NAMES.get(module_id, module_id)
            algorithm = data.get('algorithm', 'unknown')

            # 构建可读文本内容
            content = f"# 实验结果 - {module_name} (ID: {experiment_id})\n\n"

            # 算法信息
            content += "## 算法信息\n"
            content += f"- 算法: {algorithm}\n"
            content += f"- 模块: {module_name}\n\n"

            # 参数配置
            if 'parameters' in data:
                content += "## 参数配置\n"
                for key, value in data['parameters'].items():
                    content += f"- {key}: {value}\n"
                content += "\n"

            # 性能指标
            if 'metrics' in data:
                content += "## 性能指标\n"
                for key, value in data['metrics'].items():
                    if isinstance(value, float):
                        content += f"- {key}: {value:.4f}\n"
                    else:
                        content += f"- {key}: {value}\n"
                content += "\n"

            # 执行信息
            if 'execution_info' in data:
                exec_info = data['execution_info']
                content += "## 执行信息\n"
                if 'start_time' in exec_info:
                    content += f"- 开始时间: {exec_info['start_time']}\n"
                if 'end_time' in exec_info:
                    content += f"- 结束时间: {exec_info['end_time']}\n"
                if 'duration_seconds' in exec_info:
                    duration_min = exec_info['duration_seconds'] / 60
                    content += f"- 执行时长: {duration_min:.1f}分钟\n"
                content += "\n"

            # 输出文件
            if 'output_files' in data:
                content += "## 输出文件\n"
                for file in data['output_files']:
                    content += f"- {file}\n"
                content += "\n"

            metadata = {
                'source': self.file_path,
                'type': 'experiment_result',
                'filename': os.path.basename(self.file_path),
                'experiment_id': experiment_id,
                'module_id': module_id,
                'module_name': module_name,
                'algorithm': algorithm
            }

            return [Document(content=content, metadata=metadata)]

        except json.JSONDecodeError as e:
            logger.error(f"JSON解析失败 {self.file_path}: {e}")
            return []
        except Exception as e:
            logger.error(f"加载实验结果失败 {self.file_path}: {e}")
            return []


class LogFileLoader:
    """日志文件加载器"""

    def __init__(self, file_path: str):
        self.file_path = file_path

    def load(self) -> List[Document]:
        """加载日志文件，提取关键信息"""
        try:
            with open(self.file_path, 'r', encoding='utf-8', errors='ignore') as f:
                lines = f.readlines()

            if not lines:
                return []

            # 提取实验ID和模块ID（从文件名）
            filename = os.path.basename(self.file_path)
            # 假设文件名格式: module_id_experiment_id.log
            parts = filename.replace('.log', '').split('_')
            module_id = parts[0] if len(parts) > 0 else 'unknown'
            experiment_id = parts[-1] if len(parts) > 1 and parts[-1].isdigit() else 'unknown'

            # 提取错误信息
            errors = []
            for i, line in enumerate(lines):
                if any(keyword in line.upper() for keyword in ['ERROR', 'CRITICAL', 'EXCEPTION', 'FAILED', 'TIMEOUT']):
                    # 提取上下文（前后3行）
                    start = max(0, i - 3)
                    end = min(len(lines), i + 4)
                    context = ''.join(lines[start:end])

                    errors.append({
                        'line_number': i + 1,
                        'content': line.strip(),
                        'context': context
                    })

            # 提取时间范围
            first_line = lines[0].strip()
            last_line = lines[-1].strip()

            # 尝试提取时间戳（假设格式：YYYY-MM-DD HH:MM:SS）
            timestamp_pattern = r'\d{4}-\d{2}-\d{2}\s+\d{2}:\d{2}:\d{2}'
            start_time_match = re.search(timestamp_pattern, first_line)
            end_time_match = re.search(timestamp_pattern, last_line)

            start_time = start_time_match.group(0) if start_time_match else 'unknown'
            end_time = end_time_match.group(0) if end_time_match else 'unknown'

            # 构建文档内容
            content = f"# 实验日志 - {module_id} (ID: {experiment_id})\n\n"

            content += "## 执行概要\n"
            content += f"- 开始时间: {start_time}\n"
            content += f"- 结束时间: {end_time}\n"
            content += f"- 总日志条数: {len(lines)}\n"
            content += f"- 错误数量: {len(errors)}\n\n"

            # 关键事件（前10行和后10行）
            content += "## 关键事件\n"
            content += "### 开始阶段\n"
            for line in lines[:10]:
                if line.strip():
                    content += f"- {line.strip()}\n"
            content += "\n### 结束阶段\n"
            for line in lines[-10:]:
                if line.strip():
                    content += f"- {line.strip()}\n"
            content += "\n"

            # 错误详情
            if errors:
                content += "## 错误详情\n"
                for idx, error in enumerate(errors[:5], 1):  # 只取前5个错误
                    content += f"### 错误{idx}\n"
                    content += f"- 行号: {error['line_number']}\n"
                    content += f"- 内容: {error['content']}\n"
                    content += f"- 上下文:\n```\n{error['context']}```\n\n"

            metadata = {
                'source': self.file_path,
                'type': 'log',
                'filename': os.path.basename(self.file_path),
                'module_id': module_id,
                'experiment_id': experiment_id,
                'error_count': len(errors),
                'start_time': start_time,
                'end_time': end_time
            }

            return [Document(content=content, metadata=metadata)]

        except Exception as e:
            logger.error(f"加载日志文件失败 {self.file_path}: {e}")
            return []


class DocumentLoader:
    """统一文档加载器"""

    LOADERS = {
        '.md': MarkdownLoader,
        '.py': PythonLoader,
        '.js': JavaScriptLoader,
        '.docx': DocxLoader,
        '.pdf': PDFLoader,
        '.json': ExperimentResultLoader,
        '.log': LogFileLoader,
    }

    @classmethod
    def load(cls, file_path: str) -> List[Document]:
        """根据文件类型自动选择加载器"""
        ext = Path(file_path).suffix.lower()

        loader_class = cls.LOADERS.get(ext)
        if not loader_class:
            logger.warning(f"不支持的文件类型: {ext}")
            return []

        loader = loader_class(file_path)
        return loader.load()

    @classmethod
    def load_directory(cls, directory: str, pattern: str = "*") -> List[Document]:
        """加载目录下的所有文档"""
        documents = []
        path = Path(directory)

        if not path.exists():
            logger.warning(f"目录不存在: {directory}")
            return []

        for file_path in path.rglob(pattern):
            if file_path.is_file():
                docs = cls.load(str(file_path))
                documents.extend(docs)

        logger.info(f"从 {directory} 加载了 {len(documents)} 个文档")
        return documents
