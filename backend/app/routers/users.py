"""
User management API routes.
"""
from __future__ import annotations
from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File, Query
from sqlalchemy.orm import Session

from ..database import get_db
from ..models import User
from ..schemas import (
    UserCreate, UserUpdate, UserResponse,
    PasswordReset, StudentImport, ErrorResponse
)
from ..services import user as user_service
from ..dependencies import get_current_teacher, get_current_admin
from .auth import get_current_user

router = APIRouter()


@router.get("", response_model=List[UserResponse])
async def get_users(
    role: Optional[str] = Query(None, description="按角色筛选: admin, teacher, student"),
    current_user: User = Depends(get_current_teacher),
    db: Session = Depends(get_db)
):
    """
    获取用户列表。
    
    仅教师和管理员可以访问。可按角色筛选。
    """
    users = user_service.get_all_users(db, role=role)
    return users


@router.post("", response_model=UserResponse, status_code=status.HTTP_201_CREATED, responses={
    400: {"model": ErrorResponse, "description": "创建失败"}
})
async def create_user(
    user_data: UserCreate,
    current_user: User = Depends(get_current_teacher),
    db: Session = Depends(get_db)
):
    """
    创建新用户。
    
    仅教师和管理员可以创建用户。
    """
    user = user_service.create_user(db, user_data)
    return user


@router.get("/{user_id}", response_model=UserResponse, responses={
    404: {"model": ErrorResponse, "description": "用户不存在"}
})
async def get_user(
    user_id: int,
    current_user: User = Depends(get_current_teacher),
    db: Session = Depends(get_db)
):
    """获取用户详情。"""
    user = user_service.get_user_by_id(db, user_id)
    if not user:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail={
                "code": "RES_001",
                "message": "用户不存在",
                "suggestion": "请检查用户ID是否正确"
            }
        )
    return user


@router.put("/{user_id}", response_model=UserResponse, responses={
    404: {"model": ErrorResponse, "description": "用户不存在"}
})
async def update_user(
    user_id: int,
    user_data: UserUpdate,
    current_user: User = Depends(get_current_teacher),
    db: Session = Depends(get_db)
):
    """更新用户信息。"""
    user = user_service.update_user(db, user_id, user_data)
    return user


@router.put("/{user_id}/reset-password", response_model=UserResponse, responses={
    404: {"model": ErrorResponse, "description": "用户不存在"}
})
async def reset_password(
    user_id: int,
    password_data: PasswordReset,
    current_user: User = Depends(get_current_teacher),
    db: Session = Depends(get_db)
):
    """
    重置用户密码。
    
    仅教师和管理员可以重置密码。
    """
    user = user_service.reset_password(db, user_id, password_data)
    return user


@router.delete("/{user_id}", status_code=status.HTTP_204_NO_CONTENT, responses={
    404: {"model": ErrorResponse, "description": "用户不存在"}
})
async def delete_user(
    user_id: int,
    current_user: User = Depends(get_current_admin),
    db: Session = Depends(get_db)
):
    """
    删除用户。
    
    仅管理员可以删除用户。
    """
    user_service.delete_user(db, user_id)
    return None


@router.post("/import", status_code=status.HTTP_201_CREATED)
async def import_students(
    file: UploadFile = File(..., description="Excel文件 (.xlsx)"),
    current_user: User = Depends(get_current_teacher),
    db: Session = Depends(get_db)
):
    """
    从Excel文件批量导入学生。
    
    Excel文件需包含列: username, real_name, student_id, class_name(可选)。
    系统将自动生成初始密码。
    """
    # Validate file type
    if not file.filename or not file.filename.endswith('.xlsx'):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail={
                "code": "BIZ_101",
                "message": "不支持的文件格式",
                "suggestion": "请上传 .xlsx 格式的Excel文件"
            }
        )
    
    try:
        import openpyxl
        from io import BytesIO
        
        content = await file.read()
        wb = openpyxl.load_workbook(BytesIO(content))
        ws = wb.active
        
        # Read header row
        headers = [cell.value for cell in ws[1]] if ws.max_row >= 1 else []
        
        # Map column indices
        col_map = {}
        required_cols = ["username", "real_name", "student_id"]
        for i, header in enumerate(headers):
            if header in ["username", "real_name", "student_id", "class_name"]:
                col_map[header] = i
        
        # Check required columns
        for col in required_cols:
            if col not in col_map:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail={
                        "code": "VAL_001",
                        "message": f"Excel文件缺少必需列: {col}",
                        "suggestion": "请确保Excel文件包含 username, real_name, student_id 列"
                    }
                )
        
        # Parse student data
        students = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[col_map["username"]]:
                continue
            
            student = StudentImport(
                username=str(row[col_map["username"]]).strip(),
                real_name=str(row[col_map["real_name"]]).strip(),
                student_id=str(row[col_map["student_id"]]).strip(),
                class_name=str(row[col_map["class_name"]]).strip() if "class_name" in col_map and row[col_map["class_name"]] else None
            )
            students.append(student)
        
        if not students:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail={
                    "code": "VAL_001",
                    "message": "Excel文件中没有有效的学生数据",
                    "suggestion": "请检查文件内容"
                }
            )
        
        results = user_service.batch_import_students(db, students)
        
        created_count = sum(1 for r in results if r["status"] == "created")
        skipped_count = sum(1 for r in results if r["status"] == "skipped")
        
        return {
            "message": f"导入完成: 成功创建 {created_count} 个账户, 跳过 {skipped_count} 个",
            "total": len(results),
            "created": created_count,
            "skipped": skipped_count,
            "details": results
        }
    
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail={
                "code": "SRV_001",
                "message": "导入失败",
                "suggestion": f"请检查Excel文件格式是否正确: {str(e)}"
            }
        )
